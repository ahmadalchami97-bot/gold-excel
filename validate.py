#!/usr/bin/env python3
"""Independently re-implement the indicators in Python and compare against the
formula-evaluated values from the generated workbook. Validates engine math."""
import sys, math, warnings
warnings.filterwarnings("ignore")
import numpy as np
import openpyxl
import formulas
from xlsxwriter.utility import xl_col_to_name

PATH = sys.argv[1] if len(sys.argv) > 1 else "/tmp/tiny.xlsx"
BASENAME = PATH.split("/")[-1]

CALC_FIELDS = [
    "date","open","high","low","close","idx","ret","ema20","ema50","ema100","ema200",
    "gain","loss","rsi","ema12","ema26","macd","macd_sig","macd_hist","hh14","ll14",
    "stoch_k","stoch_d","rsi_min","rsi_max","srsi_raw","srsi_k","srsi_d","tr","atr",
    "plus_dm","minus_dm","plus_di","minus_di","dx","adx","bb_mid","bb_std","bb_up",
    "bb_low","bb_pctb","bb_bw","don_up","don_low","don_mid","hv","tenkan","kijun",
    "senkouA","senkouB","fib_hh","fib_ll","reg_slope","reg_int","reg_mid","reg_steyx",
    "swing_hi","swing_lo",
]
COL = {f: xl_col_to_name(i) for i, f in enumerate(CALC_FIELDS)}

# ---- read Data ----
wb = openpyxl.load_workbook(PATH)
dws = wb["Data"]
O, H, L, C = [], [], [], []
r = 2
while dws.cell(r, 5).value not in (None, ""):
    O.append(float(dws.cell(r, 2).value)); H.append(float(dws.cell(r, 3).value))
    L.append(float(dws.cell(r, 4).value)); C.append(float(dws.cell(r, 5).value))
    r += 1
O, H, L, C = map(np.array, (O, H, L, C))
n = len(C)
print(f"{n} data rows")

# ---- reference indicators ----
def ema(x, p):
    k = 2 / (p + 1); out = np.empty(len(x)); out[0] = x[0]
    for i in range(1, len(x)):
        out[i] = x[i] * k + out[i - 1] * (1 - k)
    return out

ref = {}
for p, nm in [(20,"ema20"),(50,"ema50"),(100,"ema100"),(200,"ema200"),(12,"ema12"),(26,"ema26")]:
    ref[nm] = ema(C, p)
gain = np.maximum(np.diff(C, prepend=C[0]), 0); gain[0] = 0
loss = np.maximum(-np.diff(C, prepend=C[0]), 0); loss[0] = 0
gain[0] = np.nan; loss[0] = np.nan  # first has no prev
rsi = np.full(n, np.nan)
for i in range(n):
    if i >= 14:
        ag = np.nanmean(gain[i-13:i+1]); al = np.nanmean(loss[i-13:i+1])
        rsi[i] = 100 if al == 0 else 100 - 100/(1 + ag/al)
ref["rsi"] = rsi
macd = ref["ema12"] - ref["ema26"]; ref["macd"] = macd
sig = ema(macd, 9); ref["macd_sig"] = sig; ref["macd_hist"] = macd - sig
stoch_k = np.full(n, np.nan)
for i in range(n):
    if i >= 13:
        hh = H[i-13:i+1].max(); ll = L[i-13:i+1].min()
        stoch_k[i] = np.nan if hh == ll else 100*(C[i]-ll)/(hh-ll)
ref["stoch_k"] = stoch_k
stoch_d = np.full(n, np.nan)
for i in range(n):
    if i >= 15 and not np.isnan(stoch_k[i-2:i+1]).any():
        stoch_d[i] = stoch_k[i-2:i+1].mean()
ref["stoch_d"] = stoch_d
tr = np.empty(n); tr[0] = H[0]-L[0]
for i in range(1, n):
    tr[i] = max(H[i]-L[i], abs(H[i]-C[i-1]), abs(L[i]-C[i-1]))
atr = np.full(n, np.nan)
for i in range(n):
    if i >= 14: atr[i] = tr[i-13:i+1].mean()
ref["atr"] = atr
bb_mid = np.full(n, np.nan); bb_up = np.full(n, np.nan); bb_low = np.full(n, np.nan)
for i in range(n):
    if i >= 19:
        w = C[i-19:i+1]; m = w.mean(); s = w.std()  # population std
        bb_mid[i], bb_up[i], bb_low[i] = m, m+2*s, m-2*s
ref["bb_mid"], ref["bb_up"], ref["bb_low"] = bb_mid, bb_up, bb_low
don_up = np.full(n, np.nan); don_low = np.full(n, np.nan)
for i in range(n):
    if i >= 19: don_up[i] = H[i-19:i+1].max(); don_low[i] = L[i-19:i+1].min()
ref["don_up"], ref["don_low"] = don_up, don_low
logret = np.full(n, np.nan)
for i in range(1, n): logret[i] = math.log(C[i]/C[i-1])
hv = np.full(n, np.nan)
for i in range(n):
    if i >= 20:  # need 20 returns in window i-19..i (all defined since i>=20)
        hv[i] = np.std(logret[i-19:i+1], ddof=1) * math.sqrt(252)
ref["hv"] = hv
tenkan = np.full(n, np.nan); kijun = np.full(n, np.nan)
for i in range(n):
    if i >= 8: tenkan[i] = (H[i-8:i+1].max()+L[i-8:i+1].min())/2
    if i >= 25: kijun[i] = (H[i-25:i+1].max()+L[i-25:i+1].min())/2
ref["tenkan"], ref["kijun"] = tenkan, kijun
fib_hh = np.full(n, np.nan); fib_ll = np.full(n, np.nan)
for i in range(n):
    if i >= 19:
        w0 = max(0, i-99); fib_hh[i] = H[w0:i+1].max(); fib_ll[i] = L[w0:i+1].min()
ref["fib_hh"], ref["fib_ll"] = fib_hh, fib_ll
reg_mid = np.full(n, np.nan)
idx = np.arange(1, n+1)
for i in range(n):
    if i >= 19:
        w0 = max(0, i-99); x = idx[w0:i+1]; y = C[w0:i+1]
        sl, inter = np.polyfit(x, y, 1)
        reg_mid[i] = inter + sl*idx[i]
ref["reg_mid"] = reg_mid

# ---- evaluate workbook ----
xl = formulas.ExcelModel().loads(PATH).finish()
sol = xl.calculate()

def cell(sheet, addr):
    k = f"'[{BASENAME}]{sheet.upper()}'!{addr}"
    if k not in sol: return None
    v = sol[k]
    v = getattr(v, "value", v)
    try:
        v = np.asarray(v).ravel()[0]
    except Exception:
        pass
    return v

# check the last data row + a couple of mid rows
rows = sorted(set([n+1, n, max(60, n-1), max(45, n//2)]))
checks = ["ema20","ema50","ema100","ema200","rsi","macd","macd_sig","macd_hist",
          "stoch_k","stoch_d","atr","bb_mid","bb_up","bb_low","don_up","don_low",
          "hv","tenkan","kijun","fib_hh","fib_ll","reg_mid"]
fails = 0; errs = 0; checked = 0
for wr in rows:
    di = wr - 2  # 0-based data index
    if di < 0 or di >= n: continue
    for f in checks:
        got = cell("Calc_Daily", f"{COL[f]}{wr}")
        exp = ref[f][di]
        if isinstance(got, str) and got.startswith("#"):
            print(f"  ERROR {f} row{wr}: {got}"); errs += 1; continue
        if isinstance(exp, float) and math.isnan(exp):
            # expect blank ("") in workbook
            if got not in ("", None) and not (isinstance(got, float) and math.isnan(got)):
                # tolerate tiny: it should be blank
                print(f"  NB   {f} row{wr}: expected blank, got {got!r}")
            continue
        if got in ("", None):
            print(f"  MISS {f} row{wr}: expected {exp:.4f}, got blank"); fails += 1; continue
        checked += 1
        if abs(float(got) - float(exp)) > max(1e-4, abs(exp)*1e-4):
            print(f"  FAIL {f} row{wr}: exp {exp:.5f} got {float(got):.5f}"); fails += 1

# sanity-range checks for adx / di / srsi at last row
for f, lo, hi in [("adx",0,100),("plus_di",0,100),("minus_di",0,100),
                  ("srsi_k",0,100),("srsi_d",0,100),("bb_pctb",-0.5,1.5),("dx",0,100)]:
    g = cell("Calc_Daily", f"{COL[f]}{n+1}")
    if isinstance(g, str) and g.startswith("#"):
        print(f"  ERROR {f}: {g}"); errs += 1
    elif g in ("", None):
        print(f"  note {f}: blank at last row")
    elif not (lo <= float(g) <= hi):
        print(f"  RANGE {f}: {g} not in [{lo},{hi}]"); fails += 1
    else:
        print(f"  ok   {f} = {float(g):.2f}")

# ---- weekly resample check ----
import datetime as dt
dates = []
r = 2
while dws.cell(r, 1).value not in (None, ""):
    dates.append(dws.cell(r, 1).value); r += 1
def monday(d): return d - dt.timedelta(days=d.weekday())
weeks = {}
for d, c, h, l, o in zip(dates, C, H, L, O):
    key = monday(d.date() if hasattr(d, "date") else d)
    if key not in weeks: weeks[key] = {"o": o, "h": h, "l": l, "c": c}
    else:
        weeks[key]["h"] = max(weeks[key]["h"], h); weeks[key]["l"] = min(weeks[key]["l"], l)
        weeks[key]["c"] = c
wkeys = sorted(weeks)
nW = len(wkeys)
last_wk = weeks[wkeys[-1]]
got_wc = cell("Calc_Weekly", f"E{nW+1}")
got_wh = cell("Calc_Weekly", f"C{nW+1}")
print(f"\nweeks: python={nW}  last close exp={last_wk['c']:.2f} got={got_wc}  "
      f"high exp={last_wk['h']:.2f} got={got_wh}")
if got_wc in ("", None) or abs(float(got_wc) - last_wk["c"]) > 0.01:
    print("  FAIL weekly close"); fails += 1
if got_wh in ("", None) or abs(float(got_wh) - last_wk["h"]) > 0.01:
    print("  FAIL weekly high"); fails += 1

print(f"\n=== checked {checked} numeric values | FAILS={fails} ERRORS={errs} ===")
sys.exit(1 if (fails or errs) else 0)
